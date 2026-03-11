import os
import random
import string
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import User
from app.services.email_service import send_import_credentials_email
from app.utils.password_utils import hash_password


UPLOAD_DIR = Path(__file__).resolve().parent.parent.parent / "uploads"


def _generate_random_password(length: int = 10) -> str:
    alphabet = string.ascii_letters + string.digits
    return "".join(random.choice(alphabet) for _ in range(length))


def import_users_from_excel(file_path: str, db: Session) -> dict[str, int]:
    workbook = load_workbook(filename=file_path)
    worksheet = workbook.active

    imported_count = 0
    skipped_count = 0

    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not row[1]:
            skipped_count += 1
            continue

        name = str(row[0]).strip()
        email = str(row[1]).strip().lower()
        phone = str(row[2]).strip() if len(row) > 2 and row[2] is not None else None
        address = str(row[3]).strip() if len(row) > 3 and row[3] is not None else None
        department = str(row[4]).strip() if len(row) > 4 and row[4] is not None else None

        existing_user = db.query(User).filter(User.email == email).first()
        if existing_user is not None:
            skipped_count += 1
            continue

        random_password = _generate_random_password()
        user = User(
            name=name,
            email=email,
            phone=phone,
            address=address,
            department=department,
            password_hash=hash_password(random_password),
            otp_code=None,
            is_verified=True,
            is_first_login=True,
            is_active=True,
        )

        db.add(user)
        db.commit()
        db.refresh(user)
        send_import_credentials_email(email, random_password)
        imported_count += 1

    workbook.close()
    return {
        "imported_count": imported_count,
        "skipped_count": skipped_count,
    }


def save_upload_file(filename: str, content: bytes) -> str:
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    file_path = UPLOAD_DIR / filename
    with open(file_path, "wb") as uploaded_file:
        uploaded_file.write(content)
    return str(file_path)


def remove_upload_file(file_path: str) -> None:
    if os.path.exists(file_path):
        os.remove(file_path)
